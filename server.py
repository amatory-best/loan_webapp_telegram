#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Backend для Mini App — читает Excel файлы из локальной папки.

Установка зависимостей:
    pip install flask flask-cors openpyxl

Запуск:
    python server.py
"""

import os
import re
import warnings
from datetime import datetime
from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS

# ═══════════════════════════════════════════════════════════
# MONKEY-PATCH: игнорируем баг openpyxl с extLst / PatternFill
# ═══════════════════════════════════════════════════════════
try:
    from openpyxl.descriptors.serialisable import Serialisable
    _orig_serialisable_init = Serialisable.__init__
    def _patched_serialisable_init(self, *args, **kwargs):
        for bad in ('extLst', 'theme'):
            kwargs.pop(bad, None)
        return _orig_serialisable_init(self, *args, **kwargs)
    Serialisable.__init__ = _patched_serialisable_init
except Exception:
    pass

try:
    from openpyxl.styles.fills import PatternFill, GradientFill
    _orig_pf_init = PatternFill.__init__
    _orig_gf_init = GradientFill.__init__
    def _patched_pf_init(self, *args, **kwargs):
        kwargs.pop('extLst', None)
        return _orig_pf_init(self, *args, **kwargs)
    def _patched_gf_init(self, *args, **kwargs):
        kwargs.pop('extLst', None)
        return _orig_gf_init(self, *args, **kwargs)
    PatternFill.__init__ = _patched_pf_init
    GradientFill.__init__ = _patched_gf_init
except Exception:
    pass

try:
    from openpyxl.worksheet.page import PrintPageSetup
    _orig_pps_init = PrintPageSetup.__init__
    def _patched_pps_init(self, *args, **kwargs):
        for bad in ('width', 'height'):
            kwargs.pop(bad, None)
        return _orig_pps_init(self, *args, **kwargs)
    PrintPageSetup.__init__ = _patched_pps_init
except Exception:
    pass

# Подавляем предупреждения openpyxl
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl.worksheet._reader")

from openpyxl import load_workbook

app = Flask(__name__)
CORS(app)

# ═══════════════════════════════════════════════════════════
# КОНФИГУРАЦИЯ
# ═══════════════════════════════════════════════════════════
EXCEL_FOLDER = r'C:\Users\Administrator\Dropbox\таблицы'
PORT = 5000

# Разрешённые файлы (имя без .xlsx)
ALLOWED_FILES = {
    'займись (женя)',
    'займись (марина)',
    'капуста (женя)',
    'капуста (марина)',
    'финкит (женя)',
    'финкит (марина)',
}

# Разрешённые листы
ALLOWED_SHEETS = {'2025', '2026', '2027'}

# ═══════════════════════════════════════════════════════════
# УТИЛИТЫ
# ═══════════════════════════════════════════════════════════

def safe_get(row, idx, default=None):
    if row is None or idx < 0:
        return default
    if idx < len(row):
        return row[idx]
    return default

def safe_float(val, default=0.0):
    if val is None: return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default

def safe_int(val, default=0):
    if val is None: return default
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return default

def parse_phone(val):
    if val is None: return False, ''
    s = str(val).strip().replace('+', '').replace(' ', '').replace('-', '').replace('(', '').replace(')', '')
    if s.startswith('375') and len(s) >= 12 and s.isdigit():
        return True, '+' + s
    return False, ''

def is_id_number(val):
    if val is None: return False
    s = str(val).strip()
    return 'PB' in s or 'VF' in s

def is_contract(val):
    if val is None: return False
    s = str(val).strip()
    return 'ФЛ' in s

def format_phone(phone):
    cleaned = phone.replace('+', '').replace(' ', '').replace('-', '')
    if len(cleaned) == 12 and cleaned.startswith('375'):
        return f"+375 {cleaned[3:5]} {cleaned[5:8]}-{cleaned[8:10]}-{cleaned[10:12]}"
    return phone

def get_initials(name):
    parts = name.split()
    if len(parts) >= 2:
        return (parts[0][0] + parts[1][0]).upper()
    return name[:2].upper()

def get_age_str(val):
    if val is None: return ''
    s = str(val).strip()
    m = re.search(r'\d+', s)
    if not m: return s
    n = int(m[0])
    if 11 <= n % 100 <= 14: return f"{n} лет"
    if n % 10 == 1: return f"{n} год"
    if 2 <= n % 10 <= 4: return f"{n} года"
    return f"{n} лет"

def get_status(status_text, return_date, today):
    s = str(status_text or '').lower().strip()
    if 'возвращ' in s: return 'closed', 'Закрыт'
    if 'просроч' in s: return 'overdue', 'Просрочка'
    if return_date and return_date < today: return 'overdue', 'Просрочка'
    if return_date and 0 <= (return_date - today).days <= 3: return 'warning', 'Истекает'
    return 'active', 'Активный'

def parse_excel_date(val):
    if val is None: return None
    if isinstance(val, datetime):
        return val
    s = str(val).strip()
    if not s: return None
    for fmt in ('%d.%m.%Y', '%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%m/%d/%Y'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None

def looks_like_fio(text):
    if not text or len(text) < 5:
        return False
    words = text.strip().split()
    if len(words) < 2 or len(words) > 5:
        return False
    cyrillic_words = [w for w in words if re.search(r'[А-Яа-яЁё]', w)]
    return len(cyrillic_words) >= 2

# ═══════════════════════════════════════════════════════════
# ПАРСИНГ ОДНОЙ СТРОКИ
# ═══════════════════════════════════════════════════════════

def parse_row(row, idx, today, source_file):
    fio_raw = safe_get(row, 1)
    fio = str(fio_raw).strip() if fio_raw is not None else ''
    if not fio or fio.lower() == 'заемщик' or len(fio) < 3:
        return None
    if not looks_like_fio(fio):
        return None

    age_str = get_age_str(safe_get(row, 2))
    score_raw = safe_get(row, 2)
    score = str(score_raw).strip() if score_raw is not None else ''
    phone = ''
    id_number = ''
    contract = ''

    for col_idx in [3, 4, 5]:
        v = safe_get(row, col_idx)
        if v is None:
            continue
        is_p, p = parse_phone(v)
        if is_p and not phone:
            phone = p
            continue
        if is_id_number(v) and not id_number:
            id_number = str(v).strip()
            continue
        if is_contract(v) and not contract:
            contract = str(v).strip()
            continue

    # Рейтинг (колонка R — Р.З.) и доп.информация (колонка T — доп.инфа)
    rating_raw = safe_get(row, 17)
    rating = str(rating_raw).strip() if rating_raw is not None else ''
    extra_raw = safe_get(row, 19)
    extra_info = str(extra_raw).strip() if extra_raw is not None else ''

    # Fallback телефон из колонки S (18)
    if not phone:
        v18 = safe_get(row, 18)
        if v18 is not None:
            is_p, p = parse_phone(v18)
            if is_p:
                phone = p

    issue_date = parse_excel_date(safe_get(row, 6))
    return_date = parse_excel_date(safe_get(row, 15))
    issue_date_str = issue_date.strftime('%d.%m.%Y') if issue_date else ''
    return_date_str = return_date.strftime('%d.%m.%Y') if return_date else ''

    amount = safe_float(safe_get(row, 7))
    term = safe_int(safe_get(row, 8))
    rate = safe_float(safe_get(row, 9))
    total_return = safe_float(safe_get(row, 10))
    taxes = safe_float(safe_get(row, 11))
    service_fee = safe_float(safe_get(row, 12))
    to_card = safe_float(safe_get(row, 13))
    profit = safe_float(safe_get(row, 14))

    status_code, status_label = get_status(safe_get(row, 16), return_date, today)

    days_left = 0
    if return_date:
        days_left = (return_date - today).days
    if status_code == 'closed':
        days_left = 0

    if amount <= 0 and total_return <= 0 and profit <= 0:
        return None

    return {
        "id": contract if contract else f"LOAN-{idx}",
        "amount": amount,
        "currency": "₽",
        "status": status_code,
        "statusLabel": status_label,
        "clientName": fio,
        "clientInitials": get_initials(fio),
        "term": f"{term} дн.",
        "rate": f"{rate}%/день",
        "issueDate": issue_date_str,
        "returnDate": return_date_str,
        "daysLeft": days_left,
        "profit": round(profit, 2),
        "totalReturn": round(total_return, 2),
        "finance": {
            "taxes": round(taxes, 2),
            "serviceFee": round(service_fee, 2),
            "toCard": round(to_card, 2),
            "netProfit": round(profit, 2)
        },
        "client": {
            "fullName": fio,
            "initials": get_initials(fio),
            "phone": phone,
            "phoneFormatted": format_phone(phone) if phone else '',
            "age": age_str,
            "idNumber": id_number,
            "score": score,
            "rating": rating
        },
        "contract": {
            "number": contract if contract else f"2026-ФЛ-{idx:06d}",
            "url": "#"
        },
        "source": source_file,
        "extraInfo": extra_info
    }

# ═══════════════════════════════════════════════════════════
# ЧТЕНИЕ ФАЙЛА ЧЕРЕЗ OPENPYXL (read_only — игнорирует стили)
# ═══════════════════════════════════════════════════════════

def read_excel_file(filepath):
    loans = []
    filename = os.path.basename(filepath)
    name_no_ext = os.path.splitext(filename)[0].lower()

    if name_no_ext not in ALLOWED_FILES:
        print(f"[SKIP] {filename}: не в списке разрешённых")
        return loans

    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)

    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        print(f"[ERROR] Не удалось открыть {filename}: {e}")
        return loans

    try:
        sheet_names = wb.sheetnames
    except Exception as e:
        print(f"[ERROR] Не удалось получить листы {filename}: {e}")
        wb.close()
        return loans

    for sheet_name in sheet_names:
        if sheet_name not in ALLOWED_SHEETS:
            print(f"[SKIP] {filename} / '{sheet_name}': не в списке разрешённых листов")
            continue

        try:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            start_row = 0
            for i in range(min(20, len(rows))):
                row = rows[i]
                fio_candidate = safe_get(row, 1)
                if fio_candidate is not None and looks_like_fio(str(fio_candidate)):
                    start_row = i
                    break
                fio0 = safe_get(row, 0)
                if fio0 is not None and str(fio0).lower().strip() == 'заемщик':
                    start_row = i + 1
                    break

            sheet_loans = []
            for idx in range(start_row, len(rows)):
                row = rows[idx]
                try:
                    loan = parse_row(row, idx, today, filename)
                    if loan:
                        sheet_loans.append(loan)
                except Exception:
                    continue

            if sheet_loans:
                loans.extend(sheet_loans)
                print(f"[OK] {filename} / '{sheet_name}': {len(sheet_loans)} записей")
            else:
                print(f"[WARN] {filename} / '{sheet_name}': 0 записей")

        except Exception as e:
            print(f"[WARN] {filename} / '{sheet_name}': {e}")
            continue

    wb.close()
    return loans

def scan_folder():
    all_loans = []
    file_names = []

    if not os.path.exists(EXCEL_FOLDER):
        print(f"[WARN] Папка не найдена: {EXCEL_FOLDER}")
        return all_loans, file_names

    for filename in os.listdir(EXCEL_FOLDER):
        if not filename.lower().endswith('.xlsx'):
            continue
        name_no_ext = os.path.splitext(filename)[0].lower()
        if name_no_ext not in ALLOWED_FILES:
            continue

        filepath = os.path.join(EXCEL_FOLDER, filename)
        loans = read_excel_file(filepath)
        if loans:
            all_loans.extend(loans)
            file_names.append(filename)

    return all_loans, file_names

# ═══════════════════════════════════════════════════════════
# API ENDPOINTS
# ═══════════════════════════════════════════════════════════

@app.route('/')
def index():
    return send_from_directory('.', 'index.html')

@app.route('/<path:filename>')
def static_files(filename):
    if os.path.exists(filename):
        return send_from_directory('.', filename)
    return jsonify({"error": "Not found"}), 404

@app.route('/api/loans', methods=['GET'])
def get_loans():
    loans, files = scan_folder()
    return jsonify({
        "success": True,
        "loans": loans,
        "files": files,
        "folder": EXCEL_FOLDER,
        "timestamp": datetime.now().isoformat()
    })

@app.route('/api/health', methods=['GET'])
def health():
    return jsonify({
        "status": "ok",
        "folder": EXCEL_FOLDER,
        "folder_exists": os.path.exists(EXCEL_FOLDER)
    })

@app.route('/api/config', methods=['POST'])
def set_config():
    global EXCEL_FOLDER
    data = request.get_json()
    if data and 'folder' in data:
        EXCEL_FOLDER = data['folder']
        return jsonify({"success": True, "folder": EXCEL_FOLDER})
    return jsonify({"success": False, "error": "No folder provided"}), 400

# ═══════════════════════════════════════════════════════════
# ЗАПУСК
# ═══════════════════════════════════════════════════════════

if __name__ == '__main__':
    print(f"=" * 60)
    print(f"Mini App Backend")
    print(f"Папка: {EXCEL_FOLDER}")
    print(f"Разрешённые файлы: {', '.join(ALLOWED_FILES)}")
    print(f"Разрешённые листы: {', '.join(ALLOWED_SHEETS)}")
    print(f"API: http://localhost:{PORT}/api/loans")
    print(f"=" * 60)

    if not os.path.exists(EXCEL_FOLDER):
        print(f"[WARN] Папка не существует: {EXCEL_FOLDER}")
    else:
        loans, files = scan_folder()
        print(f"[INFO] Загружено {len(loans)} записей из {len(files)} файлов")

    app.run(host='0.0.0.0', port=PORT, debug=True)
